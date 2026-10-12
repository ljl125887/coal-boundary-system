from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from collections import defaultdict

import numpy as np
from openpyxl import load_workbook
from shapely.geometry import Polygon

from .crs import lonlat_to_projected


@dataclass(frozen=True)
class ReferencePile:
    pile_id: int
    name: str
    polygon_lonlat: Polygon
    polygon_projected: Polygon
    true_volume_m3: float | None
    base_height_m: float | None


def load_reference_workbook(
    path: str | Path,
    central_meridian_deg: float = 120.0,
    false_easting_m: float = 500_000.0,
    false_northing_m: float = 0.0,
) -> list[ReferencePile]:
    wb = load_workbook(Path(path), read_only=True, data_only=True)
    ledger = wb["测量台账"]
    details = wb["顶点明细"]
    meta: dict[int, tuple[str, float | None, float | None]] = {}
    for row in ledger.iter_rows(min_row=2, values_only=True):
        if not isinstance(row[0], (int, float)):
            continue
        pile_id = int(row[0])
        meta[pile_id] = (str(row[1]), _float_or_none(row[8]), _float_or_none(row[5]))
    vertices: dict[int, list[tuple[int, float, float]]] = defaultdict(list)
    for row in details.iter_rows(min_row=2, values_only=True):
        if not isinstance(row[0], (int, float)) or row[3] is None or row[4] is None:
            continue
        vertices[int(row[0])].append((int(row[2]), float(row[3]), float(row[4])))
    result = []
    for pile_id, raw_points in sorted(vertices.items()):
        ordered = sorted(raw_points, key=lambda p: p[0])
        lon = np.asarray([p[1] for p in ordered], dtype=np.float64)
        lat = np.asarray([p[2] for p in ordered], dtype=np.float64)
        lonlat = Polygon(np.column_stack([lon, lat]))
        x, y = lonlat_to_projected(lon, lat, central_meridian_deg, false_easting_m, false_northing_m)
        projected = Polygon(np.column_stack([x, y]))
        if not lonlat.is_valid:
            lonlat = lonlat.buffer(0)
        if not projected.is_valid:
            projected = projected.buffer(0)
        name, volume, base = meta.get(pile_id, (f"煤堆#{pile_id}", None, None))
        result.append(ReferencePile(pile_id, name, lonlat, projected, volume, base))
    return result


def _float_or_none(value):
    return None if value is None else float(value)
